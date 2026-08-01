"""End-to-end exercise of every stage except the model call itself.

Run with: python -m tests.test_pipeline

Deliberately dependency-free (no pytest) so it can be run on the GPU host
straight after a clone. The model is stubbed, so this verifies the parts that
break silently — amount parsing, arithmetic checks, template alignment,
dedup — rather than the part that is obvious when it breaks.
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

import numpy as np
import yaml
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src import excel_writer  # noqa: E402
from src.preprocess import prepare  # noqa: E402
from src.schema import json_schema, load_fields  # noqa: E402
from src.store import Store, file_hash  # noqa: E402
from src.validate import _parse_amount, validate  # noqa: E402

PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASSED.append(name)
        print(f"  PASS  {name}")
    else:
        FAILED.append(f"{name}: {detail}")
        print(f"  FAIL  {name}  {detail}")


def load_cfg() -> dict:
    with (Path(__file__).resolve().parent.parent / "config.yaml").open() as fh:
        return yaml.safe_load(fh)


# ---------------------------------------------------------------------------


def test_amount_parsing() -> None:
    print("\nAmount parsing (the formats invoices actually use)")
    cases = [
        ("1,234.56", Decimal("1234.56"), "US thousands separator"),
        ("1.234,56", Decimal("1234.56"), "European notation"),
        ("$1,234.56", Decimal("1234.56"), "leading symbol"),
        ("1234.56 EUR", Decimal("1234.56"), "trailing ISO code"),
        ("(500.00)", Decimal("-500.00"), "parenthesised negative"),
        ("500.00-", Decimal("-500.00"), "trailing-minus negative"),
        ("0.00", Decimal("0.00"), "zero"),
        ("", None, "empty means absent"),
        ("N/A", None, "unparseable means absent"),
        ("1 234,56", Decimal("1234.56"), "space separator, comma decimal"),
    ]
    for raw, expected, label in cases:
        got = _parse_amount(raw)
        check(f"parse {raw!r} ({label})", got == expected, f"got {got!r}, want {expected!r}")


def test_arithmetic_validation() -> None:
    print("\nArithmetic cross-check")
    cfg = load_cfg()
    fields = load_fields(cfg)
    vcfg = cfg["validation"]

    consistent = {
        "invoice_number": "INV-001", "invoice_date": "2026-01-15", "due_date": "",
        "vendor_name": "Acme Ltd", "vendor_tax_id": "", "customer_name": "",
        "currency": "USD", "subtotal": "1000.00", "tax_amount": "200.00",
        "total_amount": "1200.00", "purchase_order": "",
    }
    result = validate(consistent, fields, vcfg)
    check("consistent invoice raises no flags", not result.flags, str(result.flags))
    check("consistent invoice scores 1.0", result.confidence == 1.0, str(result.confidence))

    broken = dict(consistent, total_amount="1500.00")
    result = validate(broken, fields, vcfg)
    check(
        "subtotal+tax != total is flagged",
        any("arithmetic" in f for f in result.flags),
        str(result.flags),
    )
    check("flagged invoice loses confidence", result.confidence < 1.0, str(result.confidence))

    rounded = dict(consistent, total_amount="1200.01")
    result = validate(rounded, fields, vcfg)
    check("1-cent rounding is tolerated", not result.flags, str(result.flags))

    missing = dict(consistent, invoice_number="", vendor_name="")
    result = validate(missing, fields, vcfg)
    check(
        "missing required fields are flagged",
        sum("required field is missing" in f for f in result.flags) == 2,
        str(result.flags),
    )

    swapped = dict(consistent, invoice_date="2026-03-01", due_date="2026-01-01")
    result = validate(swapped, fields, vcfg)
    check(
        "due date before invoice date is flagged",
        any("precedes invoice date" in f for f in result.flags),
        str(result.flags),
    )

    bad_year = dict(consistent, invoice_date="2925-01-15")
    result = validate(bad_year, fields, vcfg)
    check(
        "out-of-range year rejected (OCR digit slip)",
        any("invoice_date" in f for f in result.flags),
        str(result.flags),
    )

    symbol = dict(consistent, currency="$")
    result = validate(symbol, fields, vcfg)
    check("bare $ normalises to USD", result.values["currency"] == "USD",
          str(result.values["currency"]))


def test_schema_generation() -> None:
    print("\nGuided-decoding schema")
    cfg = load_cfg()
    fields = load_fields(cfg)
    schema = json_schema(fields)
    check("every field is required in the schema",
          len(schema["required"]) == len(fields))
    check("additionalProperties disabled",
          schema["additionalProperties"] is False)
    check("all properties are strings (absent -> empty, never invented)",
          all(p["type"] == "string" for p in schema["properties"].values()))


def test_preprocessing() -> None:
    print("\nPreprocessing (synthetic skewed invoice)")
    import cv2

    page = np.full((1200, 900, 3), 255, dtype=np.uint8)
    for i, text in enumerate(
        ["ACME LTD", "INVOICE INV-001", "Date: 2026-01-15",
         "Subtotal 1000.00", "Tax 200.00", "TOTAL 1200.00"]
    ):
        cv2.putText(page, text, (60, 150 + i * 130),
                    cv2.FONT_HERSHEY_SIMPLEX, 1.4, (0, 0, 0), 3)

    # Rotate 7 degrees and pad, simulating a phone photo on a desk.
    matrix = cv2.getRotationMatrix2D((450, 600), 7, 1.0)
    skewed = cv2.warpAffine(page, matrix, (900, 1200),
                            borderValue=(190, 190, 190))

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "skewed.png"
        cv2.imwrite(str(src), skewed)

        pages = prepare(src, load_cfg()["preprocess"])
        check("preprocessing returns one page", len(pages) == 1, str(len(pages)))
        check("output is PNG bytes", pages[0][:8] == b"\x89PNG\r\n\x1a\n")
        check("output is non-trivial in size", len(pages[0]) > 5000, str(len(pages[0])))

        # Confirm the deskewer actually reduced the skew it was given.
        from src.preprocess import _detect_skew
        restored = cv2.imdecode(np.frombuffer(pages[0], np.uint8), cv2.IMREAD_COLOR)
        residual = abs(_detect_skew(cv2.cvtColor(restored, cv2.COLOR_BGR2GRAY), 15))
        check(f"7-degree skew corrected to {residual:.2f} degrees",
              residual < 2.0, f"residual {residual:.2f}")


def test_multipage_pdf() -> None:
    print("\nMulti-page PDF rasterisation")
    import fitz

    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = Path(tmp) / "invoice.pdf"
        doc = fitz.open()
        for n in range(3):
            page = doc.new_page()
            page.insert_text((72, 100), f"Invoice page {n + 1}", fontsize=24)
        doc.save(pdf_path)
        doc.close()

        cfg = load_cfg()["preprocess"]
        pages = prepare(pdf_path, cfg)
        check("3-page PDF yields 3 images", len(pages) == 3, str(len(pages)))

        capped = prepare(pdf_path, dict(cfg, max_pages=2))
        check("max_pages caps rasterisation", len(capped) == 2, str(len(capped)))


def test_store_dedup() -> None:
    print("\nStore deduplication and idempotency")
    with tempfile.TemporaryDirectory() as tmp:
        doc_a = Path(tmp) / "a.pdf"
        doc_a.write_bytes(b"identical invoice bytes")
        doc_b = Path(tmp) / "b_renamed.pdf"
        doc_b.write_bytes(b"identical invoice bytes")

        check("same content hashes identically",
              file_hash(doc_a) == file_hash(doc_b))

        with Store(Path(tmp) / "test.db") as store:
            digest = file_hash(doc_a)
            check("unseen file is not marked processed",
                  not store.already_processed(digest))

            store.upsert(digest, "a.pdf", {"total_amount": Decimal("100.00")}, [], 1.0)
            check("file is marked processed after upsert",
                  store.already_processed(digest))

            # Re-processing the same content must update, not duplicate.
            store.upsert(digest, "b_renamed.pdf", {"total_amount": Decimal("100.00")}, [], 1.0)
            check("reprocessing updates rather than duplicates",
                  len(store.all_rows()) == 1, str(len(store.all_rows())))

            check("one row pending export", len(store.pending_export()) == 1)
            store.mark_exported([digest])
            check("nothing pending after export", len(store.pending_export()) == 0)

            store.upsert(digest, "a.pdf", {"total_amount": Decimal("999.00")}, ["x"], 0.5)
            check("re-extraction re-queues the row for export",
                  len(store.pending_export()) == 1)


def _build_template(path: Path, cfg: dict) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Register"
    headers = {
        "A": "Invoice No", "B": "Invoice Date", "C": "Due Date",
        "D": "Vendor", "E": "Vendor Tax ID", "F": "Customer",
        "G": "Currency", "H": "Subtotal", "I": "Tax", "J": "Total",
        "K": "PO Number", "L": "Source File", "M": "Confidence",
        "N": "Review Flags",
    }
    for column, title in headers.items():
        sheet[f"{column}1"] = title
    workbook.save(path)


def test_excel_export() -> None:
    print("\nExcel template export")
    cfg = load_cfg()
    fields = load_fields(cfg)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        template = tmp_path / "template.xlsx"
        _build_template(template, cfg)

        excel_cfg = dict(cfg["excel"],
                         template_path=str(template),
                         output_dir=str(tmp_path / "out"))

        warnings = excel_writer.verify_template(excel_cfg, fields)
        check("template verification passes on a matching template",
              not warnings, str(warnings))

        with Store(tmp_path / "e.db") as store:
            store.upsert(
                "hash1", "inv1.pdf",
                {
                    "invoice_number": "INV-001", "invoice_date": "2026-01-15",
                    "due_date": None, "vendor_name": "Acme Ltd",
                    "vendor_tax_id": None, "customer_name": None,
                    "currency": "USD", "subtotal": Decimal("1000.00"),
                    "tax_amount": Decimal("200.00"),
                    "total_amount": Decimal("1200.00"), "purchase_order": None,
                },
                [], 1.0,
            )
            store.upsert(
                "hash2", "inv2.pdf",
                {
                    "invoice_number": "INV-002", "invoice_date": "2026-02-01",
                    "due_date": None, "vendor_name": "Globex",
                    "vendor_tax_id": None, "customer_name": None,
                    "currency": "EUR", "subtotal": Decimal("500.00"),
                    "tax_amount": Decimal("100.00"),
                    "total_amount": Decimal("700.00"), "purchase_order": None,
                },
                ["arithmetic: subtotal 500 + tax 100 = 600, but total reads 700"],
                0.5,
            )
            rows = store.pending_export()
            out = excel_writer.export(rows, excel_cfg, fields)

        check("output file created", out.exists())
        check("template itself untouched",
              load_workbook(template).active["A2"].value is None)

        sheet = load_workbook(out).active
        check("headers preserved", sheet["A1"].value == "Invoice No", str(sheet["A1"].value))
        check("row 1 written below header", sheet["A2"].value == "INV-001", str(sheet["A2"].value))
        check("row 2 appended", sheet["A3"].value == "INV-002", str(sheet["A3"].value))
        check("amount stored as a real number, not text",
              isinstance(sheet["J2"].value, (int, float)), type(sheet["J2"].value).__name__)
        check("amount value correct", float(sheet["J2"].value) == 1200.00, str(sheet["J2"].value))
        check("date stored as a real date",
              hasattr(sheet["B2"].value, "year"), type(sheet["B2"].value).__name__)
        check("currency number format applied",
              sheet["J2"].number_format == cfg["excel"]["currency_format"],
              sheet["J2"].number_format)
        check("confidence written to audit column",
              float(sheet["M2"].value) == 1.0, str(sheet["M2"].value))
        check("review flags written for the bad row",
              "arithmetic" in str(sheet["N3"].value), str(sheet["N3"].value))
        check("clean row has empty flags", sheet["N2"].value in (None, ""), str(sheet["N2"].value))


def test_export_appends_not_overwrites() -> None:
    print("\nRepeat export appends below existing rows")
    cfg = load_cfg()
    fields = load_fields(cfg)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        template = tmp_path / "template.xlsx"
        _build_template(template, cfg)

        workbook = load_workbook(template)
        workbook.active["A2"] = "PRE-EXISTING"
        workbook.save(template)

        excel_cfg = dict(cfg["excel"],
                         template_path=str(template),
                         output_dir=str(tmp_path / "out"))

        with Store(tmp_path / "a.db") as store:
            store.upsert("h1", "x.pdf", {"invoice_number": "INV-NEW"}, [], 1.0)
            out = excel_writer.export(store.pending_export(), excel_cfg, fields)

        sheet = load_workbook(out).active
        check("existing row preserved", sheet["A2"].value == "PRE-EXISTING", str(sheet["A2"].value))
        check("new row appended below it", sheet["A3"].value == "INV-NEW", str(sheet["A3"].value))


def test_template_misconfiguration_detected() -> None:
    print("\nTemplate misconfiguration detection")
    cfg = load_cfg()
    fields = load_fields(cfg)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        template = tmp_path / "sparse.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "Invoice No"   # only one header present
        workbook.save(template)

        excel_cfg = dict(cfg["excel"], template_path=str(template))
        warnings = excel_writer.verify_template(excel_cfg, fields)
        check("missing headers produce warnings", len(warnings) > 5, str(len(warnings)))

        bad = dict(excel_cfg, columns={**cfg["excel"]["columns"], "not_a_field": "Z"})
        warnings = excel_writer.verify_template(bad, fields)
        check("unknown mapped field is reported",
              any("unknown field" in w for w in warnings), str(warnings))

    missing_cfg = dict(cfg["excel"], template_path="/nonexistent/template.xlsx")
    try:
        excel_writer.verify_template(missing_cfg, fields)
        check("missing template raises TemplateError", False, "no exception raised")
    except excel_writer.TemplateError:
        check("missing template raises TemplateError", True)


def test_ocr_line_reconstruction() -> None:
    """The CPU backend's correctness hinges on this.

    Tesseract's own reading order interleaves a two-column invoice header, so
    "Subtotal" lands in a different line from its amount. Everything
    downstream then associates the wrong value with the label. These checks
    assert the spatial regrouping puts them back together.
    """
    print("\nOCR spatial line reconstruction (CPU backend)")
    import fitz

    from src.ocr import document_to_text, page_to_text, read_page

    cfg = load_cfg()["preprocess"]

    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp) / "two_column.pdf"
        doc = fitz.open()
        page = doc.new_page()
        # Left column label, right column value — the layout that breaks
        # naive reading order.
        page.insert_text((72, 80), "ACME SUPPLIES LIMITED", fontsize=16)
        page.insert_text((400, 80), "Invoice No: INV-2026-0417", fontsize=10)
        page.insert_text((72, 300), "Subtotal", fontsize=10)
        page.insert_text((420, 300), "1,000.00", fontsize=10)
        page.insert_text((72, 320), "VAT @ 20%", fontsize=10)
        page.insert_text((420, 320), "200.00", fontsize=10)
        doc.save(pdf)
        doc.close()

        pages = prepare(pdf, cfg)
        lines = read_page(pages[0])
        text = page_to_text(lines)

        subtotal_line = next((l for l in text.splitlines() if "Subtotal" in l), "")
        vat_line = next((l for l in text.splitlines() if "VAT" in l), "")

        check("label and value land on one line",
              "1,000.00" in subtotal_line, f"got {subtotal_line!r}")
        check("second label/value pair also reunited",
              "200.00" in vat_line, f"got {vat_line!r}")
        check("column gap marked for the model",
              "|" in subtotal_line, f"got {subtotal_line!r}")
        check("values are not cross-associated",
              "200.00" not in subtotal_line, f"got {subtotal_line!r}")

        combined = document_to_text(pages)
        check("document_to_text returns non-empty text", len(combined) > 40,
              f"{len(combined)} chars")
        # OCR text must be far smaller than image-token equivalent - this is
        # the whole reason the CPU backend is viable.
        check(f"page is ~{len(combined)//4} tokens, well under image-token cost",
              len(combined) // 4 < 600, f"{len(combined)//4} tokens")

        multi = document_to_text(pages + pages)
        check("multi-page output is labelled per page",
              "--- page 2 ---" in multi, multi[:80])


def main() -> int:
    print("=" * 68)
    print("Invoice pipeline test suite")
    print("=" * 68)

    test_amount_parsing()
    test_arithmetic_validation()
    test_schema_generation()
    test_preprocessing()
    test_multipage_pdf()
    test_ocr_line_reconstruction()
    test_store_dedup()
    test_excel_export()
    test_export_appends_not_overwrites()
    test_template_misconfiguration_detected()

    print("\n" + "=" * 68)
    print(f"{len(PASSED)} passed, {len(FAILED)} failed")
    if FAILED:
        print("\nFailures:")
        for failure in FAILED:
            print(f"  - {failure}")
        return 1
    print("=" * 68)
    return 0


if __name__ == "__main__":
    sys.exit(main())
